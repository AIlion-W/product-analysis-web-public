from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


DOCX_PATH = Path(__file__).resolve().parents[1] / "public" / "产品分析助手使用说明.docx"
NEIGONG_HEADING = "1.4．内功问诊的操作步骤"
NEXT_HEADING = "2．需要准备哪些资料"


def set_paragraph(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def set_cell(cell, text: str) -> None:
    set_paragraph(cell.paragraphs[0], text)
    for paragraph in cell.paragraphs[1:]:
        set_paragraph(paragraph, "")


def single_rfonts(properties):
    fonts = properties.findall(qn("w:rFonts"))
    if fonts:
        primary = fonts[0]
        for duplicate in fonts[1:]:
            properties.remove(duplicate)
        return primary
    return properties._add_rFonts()


def remove_generated_neigong_section(document: Document) -> None:
    paragraphs = document.paragraphs
    start = next((paragraph for paragraph in paragraphs if paragraph.text == NEIGONG_HEADING), None)
    end = next((paragraph for paragraph in paragraphs if paragraph.text == NEXT_HEADING), None)
    if start is None or end is None:
        return
    element = start._p
    while element is not None and element is not end._p:
        next_element = element.getnext()
        element.getparent().remove(element)
        element = next_element


def remove_table_by_prefix(document: Document, prefix: str) -> None:
    for table in list(document.tables):
        if table.cell(0, 0).text.startswith(prefix):
            table._tbl.getparent().remove(table._tbl)


def set_monospace_paragraph(paragraph, text: str) -> None:
    set_paragraph(paragraph, text)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(9)
    paragraph.paragraph_format.left_indent = Pt(12)
    paragraph.paragraph_format.right_indent = Pt(12)
    paragraph.paragraph_format.keep_together = True

    properties = paragraph._p.get_or_add_pPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), "F4F6F8")

    for run in paragraph.runs:
        run.font.name = "Menlo"
        run.font.size = Pt(9)
        properties = run._element.get_or_add_rPr()
        fonts = single_rfonts(properties)
        fonts.set(qn("w:ascii"), "Menlo")
        fonts.set(qn("w:hAnsi"), "Menlo")
        fonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        fonts.set(qn("w:cs"), "Menlo")


def keep_table_header_with_first_row(table) -> None:
    properties = table.rows[0]._tr.get_or_add_trPr()
    header = properties.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        properties.append(header)
    header.set(qn("w:val"), "true")
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True


def copy_row_format(source_row, target_row) -> None:
    source_properties = source_row._tr.trPr
    target_properties = target_row._tr.trPr
    if target_properties is not None:
        target_row._tr.remove(target_properties)
    if source_properties is not None:
        target_row._tr.insert(0, deepcopy(source_properties))

    for source_cell, target_cell in zip(source_row.cells, target_row.cells):
        target_cell._tc.remove(target_cell._tc.tcPr)
        target_cell._tc.insert(0, deepcopy(source_cell._tc.tcPr))

        source_paragraph_properties = source_cell.paragraphs[0]._p.pPr
        target_paragraph_properties = target_cell.paragraphs[0]._p.pPr
        if target_paragraph_properties is not None:
            target_cell.paragraphs[0]._p.remove(target_paragraph_properties)
        if source_paragraph_properties is not None:
            target_cell.paragraphs[0]._p.insert(0, deepcopy(source_paragraph_properties))


def insert_neigong_section(document: Document) -> None:
    target = next(paragraph for paragraph in document.paragraphs if paragraph.text == NEXT_HEADING)
    entries = [
        (NEIGONG_HEADING, "Heading 2", False),
        ("选择「内功问诊」，先上传我方数据包；可选添加最多 3 个竞品，每个竞品使用相同的四个文件槽位。", "Normal", False),
        ("默认排序评价和时间排序评价分别上传为两个 Excel；两个评价文件分别读取第一个工作表。", "Normal", False),
        (
            "产品数据包/\n"
            "├── 评价数据表格-默认排序.xlsx\n"
            "├── 评价数据表格-时间排序.xlsx\n"
            "├── 问大家.xlsx\n"
            "└── 评价页截图.png",
            "Normal",
            True,
        ),
        ("两个评价文件必须能映射出排名、昵称、评价时间、评价类型、SKU、初评正文和追评正文 7 个字段。常见别名会自动识别，包括「旺旺号／用户昵称」、「初评时间／评价时间」、「初评／初评内容」和「追评／追评内容」。", "Normal", False),
        ("系统会规范化全角字符与首尾空格；晒图／视频、有用、追评时间等额外的非分析列只记录警告，不阻断预检。缺少必需字段或同一字段出现两个候选列时仍会阻断，避免误映射。", "Normal", False),
        ("问大家读取第一个工作表，支持两种输入：A）精确四列「昵称、时间、问题、问答」，系统按 Excel 行序补排名；B）旧版表头必须带「序号」或「排名」，并包含「问题、回答、提问时间」。", "Normal", False),
        ("单个 XLSX 压缩文件最大 10MB；系统在解压前检查单 entry、总展开大小和压缩比，异常压缩包会在完整性检查中阻断。", "Normal", False),
        ("文件名以 .~ 或 ~$ 开头的临时锁文件会被忽略。", "Normal", False),
        ("评价与问大家各最多使用前 100 条有效数据。90 条是正式诊断门槛；默认评价达到 20 条时保留前 20 六维明细，达到 90 条才形成正式六维诊断。竞品问大家不足 90 条时，只排除该竞品的正式对照。", "Normal", False),
        ("截图支持 PNG、JPEG、WebP，并在浏览器压缩；无法压缩到 560KB 以内时，截图分析降级为 unavailable。", "Normal", False),
        ("进入完整性检查，逐个核对产品归属、实际文件、工作表名称、字段映射和使用行数；确认后才可开始分批分析。", "Normal", False),
        ("主批次失败会自动重试 1 次。失败后可手动重试；只有竞品可以排除后继续，我方产品不可排除。综合报告遇到限流、服务错误或网络错误会自动重试 1 次；最终失败后可点「重试综合报告」，且不会重跑已成功的 Excel 主任务。", "Normal", False),
        ("报告通过 finding → evidence 证据链校验后，显示数据完整性、评价类型、默认前 20 六维、问大家对照、分析结论和落地清单 6 个分区。", "Normal", False),
        ("点击「下载 HTML」获得离线六区报告；点击「导出 Excel」获得 6 个固定工作表。两种导出均来自当前已校验报告，不会再次调用模型。", "Normal", False),
        ("关闭／刷新后本次问诊状态丢失。全链路分析仍使用原通用评价逻辑，不调用内功问诊的独立解析、证据图或六区导出。", "Normal", False),
    ]
    for text, style, monospace in entries:
        paragraph = target.insert_paragraph_before(text, style=style)
        if monospace:
            set_monospace_paragraph(paragraph, text)


def apply_cjk_font_fallback(document: Document) -> None:
    font_name = "Arial Unicode MS"

    def set_fonts(rfonts) -> None:
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            rfonts.set(qn(f"w:{attribute}"), font_name)
        for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            rfonts.attrib.pop(qn(f"w:{attribute}"), None)

    for style in document.styles:
        if not hasattr(style, "font"):
            continue
        style.font.name = font_name
        properties = style._element.get_or_add_rPr()
        fonts = single_rfonts(properties)
        set_fonts(fonts)

    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in document.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)

    for paragraph in paragraphs:
        if "产品数据包/" in paragraph.text or "待分析资料/" in paragraph.text:
            continue
        for run in paragraph.runs:
            run.font.name = font_name
            properties = run._element.get_or_add_rPr()
            fonts = single_rfonts(properties)
            set_fonts(fonts)


def main() -> None:
    document = Document(DOCX_PATH)
    remove_generated_neigong_section(document)
    paragraphs = document.paragraphs

    updates = {
        6: "市场分析、多竞品对比、内功问诊与全链路分析使用说明",
        7: "先看市场，再拆竞品；标准评价数据进入可追溯问诊",
        8: "版本 1.4｜2026 年 8 月",
        13: "产品分析助手提供市场、多竞品对比、详情页、内功问诊、买家秀和全链路 6 个入口。提示词与默认产品知识位于服务器端，用户不需要自行编写分析指令。",
        14: "1.1．推荐操作路径",
        15: "先选择「市场分析」，上传行业表格、平台榜单、竞品清单、价格带或市场研究资料。",
        16: "从「对标竞品筛选表」中确认最值得正面对抗的竞品。",
        17: "再选择「多竞品对比分析」，每个竞品单独建立子文件夹，可一次上传整个文件夹或 ZIP。",
        18: "上传后先检查竞品可视化分组卡片：确认竞品数量、资料类型和每个文件的所属竞品。",
        19: "若页面出现「待归类资料」，先选择所属竞品；再勾选营销文案、视觉表达、页面框架等分析维度。",
        20: "确认页面显示「可以开始横向分析」后生成方案。结果固定为两张表，可编辑并直接导出 Excel。",
        21: "1.2．全链路分析的操作步骤",
        22: "选择「全链路分析」。",
        23: "我方产品资料可以不上传；网站已内置白云山二硫化硒产品基础知识。",
        24: "将市场、竞品、详情页、评价区和买家秀资料放入 5 个标准目录，再上传 ZIP、多个文件或整个文件夹。",
        25: "检查页面显示的 5 类资料数量和暂未识别文件数量。",
        26: "点击「开始生成全链路分析」，等待 5 个专项分析和综合汇总完成。",
        27: "结果固定为两张表：第一张汇总各环节关键问题，第二张给出去重后的 P0／P1 统一执行动作。",
        28: "1.3．六种分析方式",
        31: "2.1．我方产品资料（通用模块可选）",
        32: "通用模块已内置白云山二硫化硒产品一页纸的结构化知识。若本次产品资料有更新，可补充上传；内功问诊则按 1.4 节准备标准数据包。",
        33: "建议补充最新规格、价格、活动机制、核心卖点、完整成分表、检测报告和品牌资质。",
        34: "本次上传资料优先于内置知识；两者冲突时系统会标记「信息冲突，待确认」。",
        35: "资料应尽量保留原始文件、清晰图片和可核对的数据来源。",
        36: "图片不能证明流量、转化率、销量、成交关键词或证书真实性；缺少证据的信息会列入「待确认」。",
        38: "2.2．多竞品对比分析资料",
        39: "每个竞品建立一个独立子文件夹，例如「竞品A」「竞品B」。系统会在上传后立即显示每个竞品的资料卡片。",
        40: "每个竞品可放入主副图、详情页、营销文案、活动机制、评价截图和其他成品资料；无法自动判断归属的文件需在页面手动选择。",
        41: "2.3．全链路待分析资料",
        42: "建议将资料整理成以下 5 个目录。文件夹名称正确时，内部图片可以保留原始名称。",
        45: "系统优先根据文件夹路径和文件名分类。只要路径或文件名包含可识别关键词，就会进入对应专项提示词。",
        46: "市场：识别「市场资料」「行业分析」「市场分析」「market」。竞品：识别「竞品资料」「多竞品」「竞品分析」「competitor」。",
        47: "详情页：识别「详情页」「竞品详情」「detail」「product-detail」。评价区：识别「评价区」「评论区」「追评」「问大家」「review」。",
        48: "买家秀：识别「买家秀」「buyer」「buyer-show」。",
        52: "如果文件既没有分类文件夹，也没有可识别关键词，页面会显示为「暂未识别」。全部文件都无法识别时，全链路分析不会擅自调用专项提示词。",
        60: "原通用模块只输出一张分析／诊断主表和一张执行方案表。内功问诊显示 6 个分区，并可从同一份已校验报告下载 HTML 或导出 6 个固定 Excel 工作表。",
        65: "不上传我方产品资料可以分析吗？",
        66: "通用模块可以使用内置知识；内功问诊必须分别上传我方默认排序评价 Excel 和时间排序评价 Excel，问大家或截图缺失时对应结论会降级。",
        69: "为什么出现「待归类资料」？",
        70: "这是通用多竞品分析的文件归属提示。可在页面选择竞品归属；内功问诊则为我方和每个竞品分别提供四个固定文件槽位。",
        76: "已经选择正确的分析模块。",
        77: "通用多竞品资料已确认文件归属；若使用内功问诊，我方与最多 3 个竞品的数据包已逐个核对。",
        78: "图片清晰，文字、数据图表和证据素材没有被裁切或压缩到无法阅读。",
        79: "内功问诊的默认排序评价和时间排序评价已分别上传，两个文件的字段映射与使用行数已经确认。",
        80: "若使用全链路分析，5 个标准目录和暂未识别文件已经检查。",
    }
    for index, text in updates.items():
        set_paragraph(paragraphs[index], text)

    quick_table = next(table for table in document.tables if table.cell(0, 0).text.startswith("最简用法"))
    set_cell(
        quick_table.cell(0, 0),
        "最简用法：先做「市场分析」筛出对标竞品，再做「多竞品对比分析」；需要诊断真实评价结构时，按标准数据包进入「内功问诊」。",
    )

    analysis_table = next(table for table in document.tables if table.cell(0, 0).text == "分析方式")
    while len(analysis_table.rows) < 7:
        source_row = analysis_table.rows[-1]
        added_row = analysis_table.add_row()
        copy_row_format(source_row, added_row)
    copy_row_format(analysis_table.rows[5], analysis_table.rows[6])
    analysis_rows = [
        ("分析方式", "适用场景"),
        ("市场分析", "判断品牌、价格带、人群与卖点分布，并筛出优先对标竞品。"),
        ("多竞品对比分析", "上传后按竞品可视化分组，再横向拆解营销文案、视觉表达和页面框架。"),
        ("详情页分析", "拆解竞品逐屏结构、共性打法、差异机会和我方页面方向。"),
        ("内功问诊", "上传我方与最多 3 个竞品的默认排序评价 Excel、时间排序评价 Excel、问大家 Excel 和评价页截图；确认完整性后分批分析并导出 HTML／Excel。"),
        ("买家秀专项分析", "重点分析图片类型、文案结构、前排排序和逐条替换方向。"),
        ("全链路分析", "仍用原通用评价逻辑，一次完成 5 个旧专项并汇总统一执行清单。"),
    ]
    for row, values in zip(analysis_table.rows, analysis_rows):
        set_cell(row.cells[0], values[0])
        set_cell(row.cells[1], values[1])

    remove_table_by_prefix(document, "全链路分析资料/")
    remove_table_by_prefix(document, "市场资料_平台榜单.xlsx")
    set_monospace_paragraph(
        paragraphs[43],
        "待分析资料/\n"
        "├── 01_市场资料/\n"
        "├── 02_竞品资料/\n"
        "│   ├── 竞品A/\n"
        "│   └── 竞品B/\n"
        "├── 03_详情页/\n"
        "├── 04_评价区/\n"
        "└── 05_买家秀/",
    )
    set_monospace_paragraph(
        paragraphs[51],
        "市场资料_平台榜单.xlsx\n"
        "竞品资料/竞品A/详情页_01.png\n"
        "竞品资料/竞品B/营销文案.pdf\n"
        "评价区_问大家.png\n"
        "买家秀_前排01.png",
    )

    warning_table = next(
        table for table in document.tables if table.cell(0, 0).text.startswith("注意：")
    )
    set_cell(
        warning_table.cell(0, 0),
        "注意：通用多竞品资料必须分开建子文件夹；内功问诊改用每个产品独立的固定文件槽位。",
    )

    upload_table = next(
        table for table in document.tables if table.cell(0, 0).text == "项目"
    )
    upload_updates = {
        "图片分析版": ("单次直传预算", "通用模块当前 Sites 线上总计建议不超过 850KB"),
        "单次直传预算": ("单次直传预算", "通用模块当前 Sites 线上总计建议不超过 850KB"),
        "单个直接上传文件": ("单个直接上传文件", "受 850KB 单次总预算约束，较大文件需拆分或压缩"),
        "单次上传总大小": ("单次上传总大小", "通用模块当前 Sites 线上约 850KB"),
    }
    for row in upload_table.rows:
        label = row.cells[0].text
        if label in upload_updates:
            new_label, value = upload_updates[label]
            set_cell(row.cells[0], new_label)
            set_cell(row.cells[1], value)
    keep_table_header_with_first_row(upload_table)

    insert_neigong_section(document)
    apply_cjk_font_fallback(document)
    document.save(DOCX_PATH)


if __name__ == "__main__":
    main()
